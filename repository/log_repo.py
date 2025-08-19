from motor.motor_asyncio import AsyncIOMotorClient
from repository.base import DatabaseRepository
from loguru import logger
import os
from openpyxl import Workbook

from core.config import settings
from services.clean_message import clean_log_message
from services.log_label import LogLabeler

class LogRepository(DatabaseRepository):
    def __init__(self, mongo: AsyncIOMotorClient) -> None:
        super().__init__(mongo)
        self.coll = self.get_collection()

    async def get_date(self):
        pipeline = [
            {
                '$group': {
                    '_id': {
                        '$dateToString': {
                            'format': '%Y-%m-%d',
                            'date': '$timestamp',
                        }
                    }
                }
            },
            { '$sort': { '_id': -1 } },
            { '$project': { 'dates': '$_id', '_id': 0 } }
        ]

        date_lst = self.coll.aggregate(pipeline)


        return {'dates' : [ doc['dates'] async for doc in date_lst ]}

    async def save_data(self):
        loglabeler = LogLabeler()
        os.makedirs("output", exist_ok=True)

        existing_files = {
            f.split("_")[1].split(".")[0]
            for f in os.listdir("output")
            if f.endswith(".xlsx")
        }

        logger.info(f"Found {len(existing_files)} existing files")

        all_days_pipeline = [
            {
                "$group": {
                    "_id": {
                        "$dateToString": {"format": "%Y-%m-%d", "date": "$timestamp"}
                    }
                }
            },
            {"$sort": {"_id": 1}}
        ]

        all_days = await self.coll.aggregate(all_days_pipeline).to_list()
        all_days = [data["_id"] for data in all_days]

        missing_day = [day for day in all_days if day not in existing_files]
        
        if not missing_day:
            return {}
        
        pipeline = [
            {
                "$match": {
                    "$expr": {
                        "$in": [
                            {"$dateToString": {"format": "%Y-%m-%d", "date": "$timestamp"}},
                            missing_day
                        ]
                    }
                }
            },
            {
                "$group": {
                    "_id": {"$dateToString": {"format": "%Y-%m-%d", "date": "$timestamp"}},
                    "docs": {"$push": "$$ROOT"}
                }
            },
             { "$sort": {"_id": 1} }
        ]


        datas = await self.coll.aggregate(pipeline).to_list()

        for data in datas:
            day = data['_id']
            docs = data['docs']

            file_name = f'output/log_{day}.xlsx'


            wb = Workbook()
            ws = wb.create_sheet(title='DATA')

            if "Sheet" in wb.sheetnames:
                std = wb["Sheet"]
                wb.remove(std)


            headers = ['id', 'timestamp', 'severity', 'message', 'label', 'data_extracted']

            ws.append(headers)

            for doc in docs:
                clean_msg = clean_log_message(doc.get('message', ''))
                label, data_extract = loglabeler.label_log_by_regex(clean_msg)
                ws.append([
                    str(doc["_id"]),
                    str(doc["timestamp"]),
                    doc.get("severity"),
                    doc.get("message"),
                    str(label), 
                    str(data_extract)
                ])
            
            wb.save(file_name, )
            logger.info(f"Saved {len(docs)} docs for {day}")

        return True
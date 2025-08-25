import os
from openpyxl import Workbook
from loguru import logger

from services import log_label

def export_log(datas):
    os.makedirs('output', exist_ok=True)

    for data in datas:
        day = data['_id']
        values = data['docs']

        file_name = f'output/log_{day}.xlsx'

        wb = Workbook()
        ws = wb.create_sheet(title='DATA')

        if 'Sheet' in wb.sheetnames:
            std = wb['Sheet']
            wb.remove(std)


        headers = ['id', 'timestamp', 'severity', 'message', 'label', 'data_extracted']

        ws.append(headers)

        for value in values:
            # clean_msg = clean_message.clean_log_message(value.get('message', ''))
            label, data_extracted = log_label.get_label_log(value.get('message', ''))

            ws.append([
                str(value['_id']),
                str(value['timestamp']),
                value.get('severity'),
                value.get('message'),
                str(label),
                str(data_extracted)
            ])

        logger.info('Have {} in {}', len(values), day)

        wb.save(file_name)

    return True